"""Build starter ``.docx`` templates for :class:`hwval.docs.SegDocument`.

The Flask config generator lets the user declare the fields that
should be exposed to the report templates and the extra context
(``project``, ``operator``, …) that should be available while
rendering. This module turns that declaration into two ready-to-open
Word documents:

* ``test_performed.docx`` — the *test performed* report.
* ``test_results.docx`` — the *test results* report.

Both documents contain the field names the user declared (as the
column headers of a results table) plus a ``{{ rec.<field> }}``
placeholder per field, and both expose every ``extra_context`` key as
a ``{{ <key> }}`` placeholder.

docxtpl syntax used
-------------------
* ``{{ var }}`` — variable substitution.
* ``{{ rec.<field> }}`` — one field of the current record; ``rec`` is
  the loop variable emitted by the ``{%tr ... %}`` row tag.
* ``{%tr for rec in <collection> %}`` / ``{%tr endfor %}`` — row-level
  loop. docxtpl **removes** the row that carries the tag, so the tag
  lives in a row of its own and the row(s) in between are repeated
  once per element of ``<collection>``.

Requires the optional ``hwval[docs]`` extra (``python-docx``); the
import is lazy so the YML part of the GUI keeps working without it.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Iterable, Mapping, Sequence

# The two report names, mirroring ``hwval.docs.REPORT_TEST_PERFORMED``
# and ``hwval.docs.REPORT_TEST_RESULTS``. Duplicated here on purpose:
# this module must stay importable without the ``hwval`` package (it
# is a pure document builder) and the names are part of its contract.
REPORT_TEST_PERFORMED = "test_performed"
REPORT_TEST_RESULTS = "test_results"
REPORTS: tuple[str, ...] = (REPORT_TEST_PERFORMED, REPORT_TEST_RESULTS)

# Context keys ``SegDocument`` always provides itself. An
# ``extra_context`` entry with one of these names would shadow the
# built-in value and break the generated templates (a shadowed
# ``records`` would stop the table loop from finding any rows).
RESERVED_CONTEXT_KEYS: frozenset[str] = frozenset(
    {
        "summary",
        "passed",
        "failed",
        "records",
        "fields",
        "include_passed",
        "include_failed",
        "generated_at",
        # Jinja literals: ``{{ true }}`` resolves to the boolean, never
        # to a context variable.
        "true",
        "false",
        "none",
    }
)


# --------------------------------------------------------------------------- #
# Field helpers                                                                #
# --------------------------------------------------------------------------- #


def fields_for_report(
    fields: Sequence[Mapping[str, Any]],
    report: str,
) -> list[str]:
    """Return the field names that belong to ``report``.

    ``fields`` is a list of mappings as produced by the form handler:
    each entry has a ``name`` and a ``reports`` key holding one of
    ``"all"``, ``"test_performed"`` or ``"test_results"``.
    """
    if report not in REPORTS:
        raise ValueError(
            f"unknown report {report!r}; valid: {list(REPORTS)}"
        )
    out: list[str] = []
    for entry in fields:
        name = entry.get("name") or ""
        if not name:
            continue
        scope = entry.get("reports") or "all"
        if scope in ("all", report):
            out.append(name)
    return out


# --------------------------------------------------------------------------- #
# docx building blocks                                                         #
# --------------------------------------------------------------------------- #


def _label(key: str) -> str:
    """Turn a context key into a human label (``dut_serial`` → ``Dut Serial``)."""
    return key.replace("_", " ").strip().title() or key


def _meta_paragraph(doc: Any, pairs: Iterable[tuple[str, str]]) -> None:
    """Add a paragraph made of bold ``label`` + ``value`` pairs.

    ``value`` is a placeholder (``{{ project }}``), so the rendered
    document shows the label and the value that was substituted at
    render time.
    """
    p = doc.add_paragraph()
    for i, (label, value) in enumerate(pairs):
        if i:
            p.add_run("    ")
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)


def _shade(cell: Any, hex_fill: str) -> None:
    """Apply a solid background colour to a table cell."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _repeat_as_header(row: Any) -> None:
    """Mark a table row as a header row that repeats on every page."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _style_tag_cell(cell: Any) -> None:
    """Render a docxtpl tag cell discreetly (small, grey).

    The row carrying a ``{%tr ... %}`` tag is removed when the
    template is rendered, but the user sees the template in Word
    first — so the tag is styled to be visible (never silently
    deleted by accident) without dominating the layout.
    """
    from docx.shared import Pt, RGBColor

    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def _add_record_table(
    doc: Any,
    field_names: Sequence[str],
    collection: str,
) -> None:
    """Add a table that repeats one row per element of ``collection``.

    The table has four rows:

    1. the header row — the field names (bold, shaded);
    2. the ``{%tr for %}`` tag row — removed by docxtpl;
    3. the data row — repeated once per record;
    4. the ``{%tr endfor %}`` tag row — removed by docxtpl.
    """
    from docx.shared import Pt

    n = len(field_names)
    table = doc.add_table(rows=4, cols=n)
    table.style = "Table Grid"

    # --- row 0: header ------------------------------------------------ #
    for i, name in enumerate(field_names):
        cell = table.rows[0].cells[i]
        cell.text = name
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
        _shade(cell, "F0F3F6")
    _repeat_as_header(table.rows[0])

    # --- row 1: loop tag (removed by docxtpl) -------------------------- #
    table.rows[1].cells[0].text = f"{{%tr for rec in {collection} %}}"
    _style_tag_cell(table.rows[1].cells[0])

    # --- row 2: the repeated data row ---------------------------------- #
    for i, name in enumerate(field_names):
        cell = table.rows[2].cells[i]
        cell.text = f"{{{{ rec.{name} }}}}"
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(10)

    # --- row 3: end tag (removed by docxtpl) --------------------------- #
    table.rows[3].cells[0].text = "{%tr endfor %}"
    _style_tag_cell(table.rows[3].cells[0])


# --------------------------------------------------------------------------- #
# Public API                                                                   #
# --------------------------------------------------------------------------- #


def build_template_docx(
    report: str,
    fields: Sequence[Mapping[str, Any]],
    extra_context: Mapping[str, Any] | None = None,
) -> bytes:
    """Build a starter ``.docx`` template for ``report``.

    Parameters
    ----------
    report:
        ``"test_performed"`` or ``"test_results"``.
    fields:
        The user's field declarations (mappings with ``name`` and
        ``reports``). Only the fields that belong to ``report`` end up
        in the generated document.
    extra_context:
        The user's extra-context mapping. Every key becomes a
        ``{{ <key> }}`` placeholder in the document header.

    Returns
    -------
    bytes
        The ``.docx`` file content, ready to be streamed to a browser.

    Raises
    ------
    ImportError
        If ``python-docx`` is not installed.
    ValueError
        If ``report`` is not one of the two known report names.
    """
    from docx import Document  # noqa: PLC0415 - lazy on purpose

    if report not in REPORTS:
        raise ValueError(
            f"unknown report {report!r}; valid: {list(REPORTS)}"
        )

    field_names = fields_for_report(fields, report)
    context = dict(extra_context or {})

    doc = Document()

    # --- document properties (not part of the rendered body) ---------- #
    props = doc.core_properties
    props.title = f"hwval {report} template"
    props.author = "hwval SegDocument config generator"
    props.comments = (
        "Starter template generated by the hwval SegDocument config "
        "generator. Edit the layout in Word, but keep the Jinja "
        "placeholders and the docxtpl row tags - docxtpl needs them "
        "to render the records."
    )

    # --- title + header block ------------------------------------------ #
    title = (
        "Test Performed"
        if report == REPORT_TEST_PERFORMED
        else "Test Results"
    )
    doc.add_heading(title, level=1)

    meta: list[tuple[str, str]] = [
        (_label(key), f"{{{{ {key} }}}}") for key in context
    ]
    if report == REPORT_TEST_RESULTS:
        meta.extend(
            [
                ("Total", "{{ summary.total }}"),
                ("Passed", "{{ summary.passed_count }}"),
                ("Failed", "{{ summary.failed_count }}"),
            ]
        )
    else:
        meta.append(("Tests performed", "{{ summary.total }}"))
    meta.append(("Generated", "{{ generated_at }}"))
    meta.append(("Fields", "{{ fields|join(', ') }}"))
    _meta_paragraph(doc, meta)

    doc.add_paragraph()

    if not field_names:
        # Nothing to tabulate for this report — say so instead of
        # emitting a broken (zero-column) table.
        doc.add_paragraph(
            "No fields are configured for this report. Add fields to "
            "the configuration and download the template again."
        )
        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()

    if report == REPORT_TEST_PERFORMED:
        doc.add_heading("Records", level=2)
        _add_record_table(doc, field_names, "records")
    else:
        doc.add_heading("Failed tests", level=2)
        _add_record_table(doc, field_names, "failed")
        doc.add_paragraph()
        doc.add_heading("Passed tests", level=2)
        _add_record_table(doc, field_names, "passed")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def template_filename(report: str) -> str:
    """The download file name for ``report`` (``test_results.docx``)."""
    if report not in REPORTS:
        raise ValueError(
            f"unknown report {report!r}; valid: {list(REPORTS)}"
        )
    return f"{report}.docx"
