"""Build the .docx templates used by example 08.

This script writes two starter templates that ``example.py`` loads
from disk. The templates are committed to the repo so the example
"just works" — but they're also the natural place to start
customising, so the structure is plain and easy to open in Word
or LibreOffice.

Edit the templates, re-run the example, and the rendered report
updates. Run this script to regenerate the originals after a
large rework::

    uv run python examples/08_test_id_and_yaml/build_templates.py

Jinja / docxtpl syntax used
----------------------------
* ``{{ var }}`` — variable substitution.
* ``{{ rec.<field> }}`` — record field; ``rec`` is the loop variable
  emitted by ``{%p for rec in records %}`` / ``{%p for rec in failed %}``.
* ``{%p for ... %}`` / ``{%p endfor %}`` — paragraph-level
  iteration. Each iteration becomes its own paragraph in the
  rendered document. The first and last lines (the tags) are
  stripped automatically by docxtpl.
* ``|join(', ')`` — Jinja filter used to render a list of field
  names as a single comma-separated string.

Available fields on each record
-------------------------------
* Built-in ``CheckRecord`` attributes: ``scope``, ``msg_id``,
  ``msg``, ``passed``, ``detail``, ``test_id``, ``value``,
  ``expected``.
* Anything you pass as a kwarg to ``check_value`` /
  ``check_value_in_range`` (e.g. ``operator=``) ends up in
  ``extras`` and is reachable as ``rec.operator``.

Available context keys
----------------------
* ``{{ project }}`` / ``{{ operator }}`` — from
  ``SegDocument(extra_context=...)``.
* ``{{ generated_at }}`` — ISO 8601 UTC timestamp.
* ``{{ summary.total }}`` / ``{{ summary.passed_count }}`` /
  ``{{ summary.failed_count }}`` — record counts (already filtered
  for the report that's about to render).
* ``{{ fields }}`` — list of field names for this report.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document  # type: ignore[import-not-found]


HERE = Path(__file__).resolve().parent
TEMPLATES = HERE / "templates"


def _heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _subtitle(doc: Document, *pairs: tuple[str, str]) -> None:
    """Add a subtitle paragraph made of ``(label, value)`` pairs.

    The ``value`` is rendered as a Jinja placeholder, so a
    template like ``("Project: ", "{{ project }}")`` ends up
    with a bold "Project: " prefix and a ``{{ project }}``
    placeholder.
    """
    p = doc.add_paragraph()
    for i, (label, value) in enumerate(pairs):
        if i:
            p.add_run("    ")
        p.add_run(label).bold = True
        p.add_run(value)


def _meta_paragraph(doc: Document, *pairs: tuple[str, str]) -> None:
    p = doc.add_paragraph()
    for i, (label, value) in enumerate(pairs):
        if i:
            p.add_run("    ")
        p.add_run(label).bold = True
        p.add_run(value)


def build_test_performed_template(path: Path) -> None:
    """Starter *test performed* template.

    Headings + subtitle + a list of records, one per paragraph.
    Edit the layout in Word — the placeholders stay in place.
    """
    doc = Document()
    _heading(doc, "Acme DUT v3 — Test Performed", level=1)
    _subtitle(
        doc,
        ("Project: ", "{{ project }}"),
        ("Operator: ", "{{ operator }}"),
        ("Generated: ", "{{ generated_at }}"),
    )
    doc.add_paragraph()  # blank spacer
    _meta_paragraph(
        doc,
        ("Tests performed: ", "{{ summary.total }}"),
        ("Fields: ", "{{ fields|join(', ') }}"),
    )
    doc.add_paragraph()
    _heading(doc, "Records", level=2)

    # The {%p %} / {%p endfor %} tags each live in their own
    # paragraph; docxtpl strips them and duplicates the middle
    # paragraph once per record.
    doc.add_paragraph("{%p for rec in records %}")
    doc.add_paragraph(
        "  - [{{ rec.test_id }}] "
        "[{{ rec.scope }}] {{ rec.msg_id }}: "
        "op={{ rec.operator }} "
        "value={{ rec.value }} expected={{ rec.expected }}"
    )
    doc.add_paragraph("{%p endfor %}")

    doc.save(path)


def build_test_results_template(path: Path) -> None:
    """Starter *test results* template — only failing records."""
    doc = Document()
    _heading(doc, "Acme DUT v3 — Test Results", level=1)
    _subtitle(
        doc,
        ("Project: ", "{{ project }}"),
        ("Operator: ", "{{ operator }}"),
    )
    doc.add_paragraph()
    _meta_paragraph(
        doc,
        ("Total: ", "{{ summary.total }}"),
        ("Passed: ", "{{ summary.passed_count }}"),
        ("Failed: ", "{{ summary.failed_count }}"),
        ("Fields: ", "{{ fields|join(', ') }}"),
    )
    doc.add_paragraph()
    _heading(doc, "Failed tests", level=2)

    doc.add_paragraph("{%p for rec in failed %}")
    doc.add_paragraph(
        "  - [{{ rec.test_id }}] op={{ rec.operator }}: "
        "{{ rec.msg }}"
    )
    doc.add_paragraph(
        "    value={{ rec.value }} expected={{ rec.expected }}"
    )
    doc.add_paragraph("    detail: {{ rec.detail }}")
    doc.add_paragraph("{%p endfor %}")

    doc.save(path)


def main() -> None:
    TEMPLATES.mkdir(parents=True, exist_ok=True)
    perf = TEMPLATES / "test_performed.docx"
    res = TEMPLATES / "test_results.docx"
    build_test_performed_template(perf)
    build_test_results_template(res)
    print(f"wrote: {perf}")
    print(f"wrote: {res}")


if __name__ == "__main__":
    main()
